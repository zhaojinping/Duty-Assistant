# -*- coding: utf-8 -*-
"""M0 核验：用真模板重跑提取，与仓库内结构快照逐项比对，并落一份脱敏结构元信息。

用法：
    python scripts/m0_verify_snapshot.py <模板目录> [--snapshot <快照JSON>] [--extra-out <输出JSON>]

- 模板目录与各输出路径一律从 CLI 读取，不硬编码机器路径（AGENTS.md）。
- 只读模板；`--extra-out` 落盘的只有结构元信息（脚注原文/数据有效性/批注/隐藏行列/序号预填），
  不含业务数据；若在数据行检出任何非预填内容，脚本拒绝落盘并以非零退出（防误传真实数据）。
- 差异数 > 0 时非零退出。
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SNAPSHOT = ROOT / "scripts" / "m0" / "excel_struct.json"

# 10 份在范围内的模板（消防器材、保护定值压板两份已停用，不在快照内）。
# 与 scripts/m0_extract_struct.py、scripts/m0_extract_extra.py 的清单保持一致。
TEMPLATES = [
    ("breaker_trip_record", "断路器跳闸记录簿.xlsx"),
    ("surge_arrester_action_record", "避雷器动作记录簿.xlsx"),
    ("grounding_wire_record", "接地线装拆记录簿.xlsx"),
    ("two_ticket_ledger", "两票登记记录簿.xlsx"),
    ("infrared_thermography_record", "设备测温记录簿.xlsx"),
    ("insulation_test_record", "绝缘测试记录簿.xlsx"),
    ("battery_voltage_test", "蓄电池电压测试记录簿.xlsx"),
    ("transformer_core_clamp_current_record", "主变铁芯夹件电流测试记录簿.xlsx"),
    ("protection_switch_record", "保护投退记录簿.xlsx"),
    ("rodent_proof_check_record", "防小动物检查记录簿.xlsx"),
]

SNAPSHOT_ROWS = 12          # 快照只抓前 12 行（见 m0_extract_struct.py）
FOOTER_SCAN_FROM = 13       # 第 13 行起为快照未覆盖区
ALLOWED_NUMERIC_PREFILL = re.compile(r"^\d+$")


def cellstr(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def full_width_merged_rows(ws):
    """整行合并（左起第 1 列到最右列）的行号，用于定位标题行与脚注行。"""
    rows = []
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 1 and rng.max_col == ws.max_column and rng.min_row == rng.max_row:
            rows.append(rng.min_row)
    return sorted(rows)


def extract_full(path):
    """真模板全行提取（结构，不落值到产物）。"""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column), start=1):
        rows.append({"r": idx, "vals": [cellstr(c.value) for c in row]})
    return {
        "sheet": ws.title,
        "dims": ws.dimensions,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "merged": sorted(str(r) for r in ws.merged_cells.ranges),
        "full_width_merged_rows": full_width_merged_rows(ws),
        "rows": rows,
    }


def extract_extra(path, sheet_struct):
    """快照未覆盖区（第 13 行起）：脚注原文、空白性核对、下拉、批注、隐藏行列、序号预填。

    第 1–12 行已由快照覆盖并与真模板逐格比对（见 compare()），此处只处理快照覆盖不到的部分，
    因此表单式的蓄电池表（条目表头在第 4 行）不会把表头误判为数据。
    """
    wb = load_workbook(path, data_only=False)
    ws = wb.worksheets[0]
    full_width = sheet_struct["full_width_merged_rows"]
    footer_row = full_width[-1] if len(full_width) > 1 else None
    scan_to = (footer_row - 1) if footer_row else ws.max_row

    # 序号预填：全表 A 列为纯数字的单元格（9 份台账式 A 列为时间/空，无此项）
    prefill_cells = [
        (cell.row, cellstr(cell.value))
        for row in ws.iter_rows(min_col=1, max_col=1)
        for cell in row
        if cellstr(cell.value) and ALLOWED_NUMERIC_PREFILL.match(cellstr(cell.value))
    ]
    prefill_rows = {r for r, _ in prefill_cells}

    footer_lines, unexpected = [], []
    for row in ws.iter_rows(min_row=FOOTER_SCAN_FROM, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            text = cellstr(cell.value)
            if text is None:
                continue
            if cell.column == 1 and cell.row in prefill_rows:
                continue
            if cell.row == footer_row:
                footer_lines.append(f"{cell.coordinate}: {text}")
            else:
                unexpected.append(f"{cell.coordinate}: {text}")

    extra = {
        "sheet": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "title_row": full_width[0] if full_width else None,
        "footer_row": footer_row,
        "footer_range": next((m for m in sheet_struct["merged"] if m.startswith(f"A{footer_row}:")), None),
        "footer_lines": footer_lines,
        "uncovered_region": {
            "from": FOOTER_SCAN_FROM,
            "to": scan_to,
            "unexpected_cells": unexpected,
        },
        "data_validations": [
            {"type": dv.type, "formula1": dv.formula1, "formula2": dv.formula2, "ranges": str(dv.sqref)}
            for dv in ws.data_validations.dataValidation
        ],
        "comments": [
            f"{c.coordinate}: {' '.join(c.comment.text.split())}"
            for row in ws.iter_rows() for c in row
            if c.comment and c.comment.text and c.comment.text.strip()
        ],
        "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden],
        "hidden_cols": [k for k, d in ws.column_dimensions.items() if d.hidden],
    }
    if prefill_cells:
        extra["prefill"] = {
            "column": "A",
            "rows": f"{prefill_cells[0][0]}-{prefill_cells[-1][0]}",
            "values": [v for _, v in prefill_cells],
        }
    return extra


def compare(snapshot, fresh, record_type):
    """逐项比对仓库快照与真模板重跑结果。"""
    diffs = []
    for key in ("sheet", "dims", "max_row", "max_col"):
        if snapshot[key] != fresh[key]:
            diffs.append(f"{key}: 快照={snapshot[key]} 真模板={fresh[key]}")
    if sorted(snapshot["merged"]) != fresh["merged"]:
        diffs.append(f"merged: 快照={sorted(snapshot['merged'])} 真模板={fresh['merged']}")
    if len(snapshot["rows"]) != SNAPSHOT_ROWS:
        diffs.append(f"快照行数 {len(snapshot['rows'])} != {SNAPSHOT_ROWS}")
    for snap_row, fresh_row in zip(snapshot["rows"], fresh["rows"][: SNAPSHOT_ROWS]):
        if snap_row["r"] != fresh_row["r"] or snap_row["vals"] != fresh_row["vals"]:
            diffs.append(f"快照第 {snap_row['r']} 行与真模板不一致")
    return diffs


def main(argv=None):
    parser = argparse.ArgumentParser(description="M0 快照 vs 真模板核验")
    parser.add_argument("template_dir", help="10 份模板所在目录")
    parser.add_argument("--snapshot", default=str(DEFAULT_SNAPSHOT), help="仓库内结构快照 JSON")
    parser.add_argument("--extra-out", default=None, help="脱敏结构元信息输出路径（可选）")
    args = parser.parse_args(argv)

    src = Path(args.template_dir)
    snapshot = json.loads(Path(args.snapshot).read_text(encoding="utf-8"))

    all_diffs, extras, missing = [], {}, []
    print(f"模板目录：{src}")
    for record_type, filename in TEMPLATES:
        path = src / filename
        if not path.exists():
            missing.append(f"{record_type}: 缺文件 {filename}")
            continue
        sheet_struct = extract_full(path)
        extras[record_type] = extract_extra(path, sheet_struct)
        snap = snapshot.get(record_type)
        if not snap:
            all_diffs.append(f"{record_type}: 快照内无该类记录")
            continue
        diffs = compare(snap["sheets"][0], sheet_struct, record_type)
        all_diffs.extend(f"{record_type}: {d}" for d in diffs)
        status = "一致" if not diffs else "★有差异"
        extra = extras[record_type]
        print(f"{record_type:<42} sheet/dims/合并/前{SNAPSHOT_ROWS}行 {status}"
              f"  未覆盖区 {extra['uncovered_region']['from']}-{extra['uncovered_region']['to']}"
              f" 异常内容 {len(extra['uncovered_region']['unexpected_cells'])} 项"
              f" | 下拉 {len(extra['data_validations'])} 批注 {len(extra['comments'])}"
              f" 隐藏行 {len(extra['hidden_rows'])} 隐藏列 {len(extra['hidden_cols'])}")

    all_diffs.extend(missing)
    print(f"\n差异合计：{len(all_diffs)}")
    for item in all_diffs:
        print("  -", item)

    if args.extra_out:
        leaked = [
            f"{rt}: {cell}" for rt, extra in extras.items()
            for cell in extra["uncovered_region"]["unexpected_cells"]
        ]
        if leaked:
            print("\n拒绝落盘：快照未覆盖区检出非脚注/非序号预填内容（疑似真实数据），请人工确认后再提交：")
            for item in leaked[:20]:
                print("  -", item)
            return 3
        payload = {
            "note": "脱敏结构元信息：脚注原文/数据有效性/批注/隐藏行列/序号预填/数据行空白性核对；不含业务数据",
            "generated_by": "scripts/m0_verify_snapshot.py",
            "records": extras,
        }
        out = Path(args.extra_out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        print(f"已写 {out}")

    return 1 if all_diffs else 0


if __name__ == "__main__":
    sys.exit(main())
